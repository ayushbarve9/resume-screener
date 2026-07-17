"""Word document (.docx) and Markdown (.md) report generator.
"""

from io import BytesIO
from datetime import datetime
from docx import Document
from docx.shared import RGBColor, Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from models import ScreeningResponse

def generate_docx_report(response: ScreeningResponse, job_title: str) -> BytesIO:
    """Generate a highly polished Word Document summarizing the screening response."""
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Style colors
    PRIMARY_RGB = RGBColor(37, 99, 235)    # #2563eb Blue
    DARK_RGB = RGBColor(15, 23, 42)        # #0f172a Slate
    GRAY_RGB = RGBColor(100, 116, 139)     # #64748b Gray
    SUCCESS_RGB = RGBColor(22, 163, 74)    # #16a34a Green
    WARNING_RGB = RGBColor(217, 119, 6)    # #d97706 Amber
    DANGER_RGB = RGBColor(220, 38, 38)     # #dc2626 Red

    # 1. Header Block
    title = doc.add_paragraph()
    title_run = title.add_run("ATS Candidate Evaluation & Recruitment Report")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = PRIMARY_RGB
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(2)

    # Subtitle with metadata
    meta = doc.add_paragraph()
    meta_run = meta.add_run(
        f"Target Role: {job_title.strip() or 'Technical Role'}  |  "
        f"Generated: {datetime.now().strftime('%b %d, %y at %I:%M %p')}  |  "
        f"Candidates Evaluated: {len(response.evaluations)}"
    )
    meta_run.font.name = "Arial"
    meta_run.font.size = Pt(10)
    meta_run.font.italic = True
    meta_run.font.color.rgb = GRAY_RGB
    meta.paragraph_format.space_after = Pt(24)

    # Divider line
    p_div = doc.add_paragraph()
    p_div_run = p_div.add_run("―" * 60)
    p_div_run.font.color.rgb = GRAY_RGB
    p_div.paragraph_format.space_after = Pt(18)

    # 2. Executive Summary (Verdicts and rankings)
    h_exec = doc.add_paragraph()
    h_exec_run = h_exec.add_run("🏆 Executive Selection Verdict")
    h_exec_run.font.name = "Arial"
    h_exec_run.font.size = Pt(15)
    h_exec_run.font.bold = True
    h_exec_run.font.color.rgb = DARK_RGB
    h_exec.paragraph_format.space_after = Pt(8)

    # Highlight winner
    p_winner = doc.add_paragraph()
    p_winner.add_run("Top Pick: ").bold = True
    p_winner_run = p_winner.add_run(response.executive_summary.winner_name or "N/A (Single Candidate)")
    p_winner_run.font.color.rgb = SUCCESS_RGB
    p_winner_run.font.bold = True
    p_winner.paragraph_format.space_after = Pt(4)

    p_just = doc.add_paragraph()
    p_just.add_run("Justification: ").bold = True
    p_just.add_run(response.executive_summary.justification)
    p_just.paragraph_format.space_after = Pt(12)

    # Trade-offs Comparison List
    if response.executive_summary.trade_offs:
        p_to_hdr = doc.add_paragraph()
        p_to_hdr.add_run("Key Comparative Trade-offs:").bold = True
        p_to_hdr.paragraph_format.space_after = Pt(4)
        for trade_off in response.executive_summary.trade_offs:
            p_to = doc.add_paragraph(style='List Bullet')
            p_to.add_run(trade_off)
            p_to.paragraph_format.space_after = Pt(3)

    # Spacing
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Ranking Table
    h_table = doc.add_paragraph()
    h_table_run = h_table.add_run("📊 Candidate Ranking Scoreboard")
    h_table_run.font.name = "Arial"
    h_table_run.font.size = Pt(14)
    h_table_run.font.bold = True
    h_table_run.font.color.rgb = DARK_RGB
    h_table.paragraph_format.space_after = Pt(8)

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Light Shading Accent 1'

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Rank'
    hdr_cells[1].text = 'Candidate Name'
    hdr_cells[2].text = 'Match Score'

    for i, name in enumerate(response.executive_summary.rankings):
        score = "N/A"
        for eval_item in response.evaluations:
            if eval_item.candidate_name.lower() == name.lower():
                score = f"{eval_item.match_score}/100"
                break
        
        row_cells = table.add_row().cells
        row_cells[0].text = str(i + 1)
        row_cells[1].text = name
        row_cells[2].text = score

    doc.add_paragraph().paragraph_format.space_after = Pt(24)

    # 3. Individual Candidate Reports
    h_reports = doc.add_paragraph()
    h_reports_run = h_reports.add_run("👤 Detailed Candidate Profile Assessments")
    h_reports_run.font.name = "Arial"
    h_reports_run.font.size = Pt(16)
    h_reports_run.font.bold = True
    h_reports_run.font.color.rgb = PRIMARY_RGB
    h_reports.paragraph_format.space_after = Pt(12)

    for eval_item in response.evaluations:
        h_cand = doc.add_paragraph()
        h_cand_run = h_cand.add_run(f"Candidate: {eval_item.candidate_name}")
        h_cand_run.font.name = "Arial"
        h_cand_run.font.size = Pt(14)
        h_cand_run.font.bold = True
        h_cand_run.font.color.rgb = DARK_RGB
        h_cand.paragraph_format.space_after = Pt(4)

        p_score = doc.add_paragraph()
        p_score.add_run("ATS Compatibility Match Score: ").bold = True
        score_val = eval_item.match_score
        score_run = p_score.add_run(f"{score_val}/100")
        score_run.font.bold = True
        if score_val >= 80:
            score_run.font.color.rgb = SUCCESS_RGB
        elif score_val >= 60:
            score_run.font.color.rgb = WARNING_RGB
        else:
            score_run.font.color.rgb = DANGER_RGB
        p_score.paragraph_format.space_after = Pt(6)

        p_gap = doc.add_paragraph()
        p_gap.add_run("Evaluation Summary: ").bold = True
        p_gap.add_run(eval_item.gap_analysis)
        p_gap.paragraph_format.space_after = Pt(10)

        # Strengths
        p_str_hdr = doc.add_paragraph()
        p_str_hdr_run = p_str_hdr.add_run("✅ Top Strengths & Match Drivers")
        p_str_hdr_run.font.bold = True
        p_str_hdr_run.font.color.rgb = SUCCESS_RGB
        p_str_hdr.paragraph_format.space_after = Pt(4)
        for strength in eval_item.strengths:
            p_s = doc.add_paragraph(style='List Bullet')
            p_s.add_run(strength)
            p_s.paragraph_format.space_after = Pt(3)

        # Gaps
        p_gap_hdr = doc.add_paragraph()
        p_gap_hdr_run = p_gap_hdr.add_run("⚠️ Identified Critical Gaps")
        p_gap_hdr_run.font.bold = True
        p_gap_hdr_run.font.color.rgb = WARNING_RGB
        p_gap_hdr.paragraph_format.space_after = Pt(4)
        for gap in eval_item.gaps:
            p_g = doc.add_paragraph(style='List Bullet')
            p_g.add_run(gap)
            p_g.paragraph_format.space_after = Pt(3)

        # Interview Questions
        p_q_hdr = doc.add_paragraph()
        p_q_hdr_run = p_q_hdr.add_run("🗣️ Target Verification Interview Questions")
        p_q_hdr_run.font.bold = True
        p_q_hdr_run.font.color.rgb = DARK_RGB
        p_q_hdr.paragraph_format.space_after = Pt(4)
        for q in eval_item.interview_questions:
            p_q = doc.add_paragraph(style='List Bullet')
            p_q.add_run(q)
            p_q.paragraph_format.space_after = Pt(3)

        # Actionable revisions
        p_rev_hdr = doc.add_paragraph()
        p_rev_hdr_run = p_rev_hdr.add_run("🛠️ Recommended Actionable Resume Revisions")
        p_rev_hdr_run.font.bold = True
        p_rev_hdr_run.font.color.rgb = PRIMARY_RGB
        p_rev_hdr.paragraph_format.space_after = Pt(4)
        for rev in eval_item.resume_revisions:
            p_r = doc.add_paragraph(style='List Bullet')
            p_r.add_run(rev)
            p_r.paragraph_format.space_after = Pt(3)

        # Divider between candidates
        p_cand_div = doc.add_paragraph()
        p_cand_div_run = p_cand_div.add_run("―" * 50)
        p_cand_div_run.font.color.rgb = GRAY_RGB
        p_cand_div.paragraph_format.space_after = Pt(18)

    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio

def generate_markdown_report(response: ScreeningResponse) -> str:
    """Generate clean GFM-compliant markdown report text."""
    md = []
    md.append("# Candidate Recruitment Assessment Report")
    md.append(f"*Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}*\n")
    
    md.append("## 🏆 Executive Verdict & Selection Summaries\n")
    md.append(f"**Top Recommended Candidate**: `{response.executive_summary.winner_name}`\n")
    md.append(f"**Justification**: {response.executive_summary.justification}\n")
    
    if response.executive_summary.trade_offs:
        md.append("### Key Trade-offs & Comparisons")
        for trade_off in response.executive_summary.trade_offs:
            md.append(f"- {trade_off}")
        md.append("")
        
    md.append("### Ranking Scoreboard")
    md.append("| Rank | Candidate Name | ATS Match Score |")
    md.append("|---|---|---|")
    for idx, name in enumerate(response.executive_summary.rankings):
        score = "N/A"
        for item in response.evaluations:
            if item.candidate_name.lower() == name.lower():
                score = f"{item.match_score}/100"
                break
        md.append(f"| {idx + 1} | {name} | {score} |")
    md.append("\n---\n")
    
    md.append("## 👤 Detailed Candidate Profiles")
    for cand in response.evaluations:
        md.append(f"### Candidate: {cand.candidate_name}")
        md.append(f"- **Target Position**: {cand.target_position}")
        md.append(f"- **Match Score**: `{cand.match_score}/100`\n")
        md.append(f"**Evaluation Summary**: {cand.gap_analysis}\n")
        
        md.append("#### Strengths")
        for strength in cand.strengths:
            md.append(f"- {strength}")
        md.append("")
        
        md.append("#### Identified Gaps")
        for gap in cand.gaps:
            md.append(f"- {gap}")
        md.append("")
        
        md.append("#### Targeted Interview Questions")
        for q in cand.interview_questions:
            md.append(f"- {q}")
        md.append("")
        
        md.append("#### Recommended Revisions")
        for rev in cand.resume_revisions:
            md.append(f"- {rev}")
        md.append("\n---\n")
        
    return "\n".join(md)
